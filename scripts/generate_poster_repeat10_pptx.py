#!/usr/bin/env python3
"""Generate an editable portrait PPTX academic poster in French."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
PLOTS = ROOT / "data" / "plots"

OUTPUT = DOCS / "poster_repeat10_fr.pptx"

NAVY = RGBColor(16, 42, 67)
BLUE = RGBColor(31, 90, 145)
BLUE_DARK = RGBColor(21, 62, 117)
TEXT = RGBColor(36, 55, 70)
TEXT_MUTED = RGBColor(92, 111, 126)
WHITE = RGBColor(255, 255, 255)
BG = RGBColor(246, 247, 245)
BORDER = RGBColor(199, 208, 216)
CHIP = RGBColor(238, 244, 250)
CHIP_ALT = RGBColor(238, 246, 245)


def add_box(slide, x, y, w, h, fill, line=BORDER, radius=True):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def add_text(slide, x, y, w, h, lines, *, font_size=18, color=TEXT, bold=False,
             font_name="Aptos", align=PP_ALIGN.LEFT, margin=0.04, spacing=1.15):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = MSO_ANCHOR.TOP

    for idx, line in enumerate(lines):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = line if line else " "
        p.alignment = align
        p.space_after = Pt(0)
        p.line_spacing = spacing
        font = p.runs[0].font
        font.name = font_name
        font.size = Pt(font_size)
        font.bold = bold
        font.color.rgb = color
    return box


def add_bullets(slide, x, y, w, h, items, *, font_size=14, color=TEXT, level=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for idx, item in enumerate(items):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = item
        p.level = level
        p.bullet = True
        p.space_after = Pt(2)
        p.line_spacing = 1.12
        font = p.runs[0].font
        font.name = "Aptos"
        font.size = Pt(font_size)
        font.color.rgb = color
    return box


def add_section(slide, x, y, w, h, title):
    add_box(slide, x, y, w, h, WHITE)
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.34)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.color.rgb = BLUE
    add_text(slide, x + 0.12, y + 0.02, w - 0.24, 0.26, [title], font_size=11, color=WHITE, bold=True)


def add_metric(slide, x, y, w, h, label, value, note, *, fill=CHIP):
    add_box(slide, x, y, w, h, fill, line=RGBColor(191, 208, 226))
    add_text(slide, x + 0.08, y + 0.07, w - 0.16, 0.16, [label], font_size=8, color=TEXT_MUTED, bold=True)
    add_text(slide, x + 0.08, y + 0.28, w - 0.16, 0.25, [value], font_size=15, color=TEXT, bold=True)
    if note:
        add_text(slide, x + 0.08, y + 0.55, w - 0.16, 0.15, [note], font_size=8, color=TEXT_MUTED)


def add_numbered(slide, x, y, w, items, *, start=1):
    cy = y
    for offset, item in enumerate(items, start=start):
        circle = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(cy), Inches(0.28), Inches(0.28)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = CHIP
        circle.line.color.rgb = BLUE
        add_text(slide, x + 0.06, cy + 0.02, 0.16, 0.16, [str(offset)], font_size=11, color=BLUE, bold=True)
        add_text(slide, x + 0.38, cy - 0.01, w - 0.38, 0.36, [item], font_size=13, color=TEXT)
        cy += 0.48


def add_picture(slide, path, x, y, w, h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def build_poster() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(33.11)
    prs.slide_height = Inches(46.81)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG

    # Header
    header = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(33.11), Inches(5.45))
    header.fill.solid()
    header.fill.fore_color.rgb = NAVY
    header.line.color.rgb = NAVY
    add_text(slide, 1.0, 0.65, 19.8, 0.35, ["TER • SDN • JUMEAU NUMÉRIQUE"], font_size=11, color=RGBColor(126, 165, 202), bold=True)
    add_text(slide, 1.0, 1.3, 19.8, 1.0, ["Jumeau numérique de réseau pour SDN :"], font_name="Georgia", font_size=24, color=WHITE, bold=True)
    add_text(slide, 1.0, 2.25, 19.8, 1.0, ["validation expérimentale après extension à REPEAT=10"], font_name="Georgia", font_size=24, color=WHITE, bold=True)
    add_text(
        slide,
        1.0,
        3.45,
        18.8,
        0.75,
        ["Prototype sur Mininet, Open vSwitch et Ryu, avec synchronisation d'état, différentiel temporel et génération d'événements."],
        font_size=12,
        color=RGBColor(215, 228, 241),
    )
    add_text(slide, 23.1, 0.95, 8.0, 0.25, ["QUESTION"], font_size=9, color=RGBColor(126, 165, 202), bold=True)
    add_text(slide, 23.1, 1.45, 8.0, 0.9, ["Le jumeau détecte-t-il les changements"], font_size=13, color=WHITE, bold=True)
    add_text(slide, 23.1, 1.95, 8.0, 0.9, ["de réseau de façon fiable et explicable ?"], font_size=13, color=WHITE, bold=True)
    add_text(slide, 23.1, 3.2, 8.0, 0.25, ["MISE À JOUR CLÉ"], font_size=9, color=RGBColor(126, 165, 202), bold=True)
    add_text(slide, 23.1, 3.65, 8.2, 0.7, ["Protocole principal porté de repeat=3 à repeat=10"], font_size=13, color=WHITE, bold=True)

    margin_x = 0.95
    gap = 0.8
    col_w = 10.2
    x1 = margin_x
    x2 = x1 + col_w + gap
    x3 = x2 + col_w + gap

    # Column 1
    add_section(slide, x1, 6.25, col_w, 5.25, "PROBLÉMATIQUE ET OBJECTIF")
    add_text(
        slide, x1 + 0.22, 6.85, col_w - 0.44, 1.65,
        [
            "Le projet ne vise pas un simple tableau de bord réseau.",
            "Il cherche à maintenir une image structurée du réseau,",
            "mise à jour via le contrôleur puis comparée dans le temps",
            "pour produire des événements interprétables."
        ],
        font_size=13,
    )
    add_bullets(
        slide, x1 + 0.22, 8.85, col_w - 0.44, 1.9,
        [
            "détecter panne et reprise de lien",
            "mesurer délai de détection et délai de reprise",
            "estimer l'impact du trafic de fond et du polling",
        ],
        font_size=13,
    )

    add_section(slide, x1, 11.9, col_w, 6.2, "CHAÎNE DU JUMEAU NUMÉRIQUE")
    add_numbered(
        slide, x1 + 0.24, 12.55, col_w - 0.5,
        [
            "Émulation du réseau dans Mininet et Open vSwitch",
            "Observation temps réel via l'API REST de Ryu",
            "Construction d'un état réseau structuré",
            "Différentiel temporel et émission d'événements",
        ],
    )
    add_text(
        slide, x1 + 0.22, 16.6, col_w - 0.44, 1.0,
        [
            "Valeur scientifique : le système apporte une sémantique",
            "historique du réseau, et non une simple vue instantanée."
        ],
        font_size=11,
        color=TEXT_MUTED,
    )

    add_section(slide, x1, 18.55, col_w, 6.75, "TOPOLOGIE ET PROTOCOLE")
    add_picture(slide, PLOTS / "report" / "topology_geant8.png", x1 + 0.22, 19.15, col_w - 0.44, 2.95)
    add_text(
        slide, x1 + 0.22, 22.3, col_w - 0.44, 0.45,
        ["Topologie principale : `geant_backbone_8cities_stage3` (8 villes, 10 liens)."],
        font_size=10,
        color=TEXT_MUTED,
    )
    add_bullets(
        slide, x1 + 0.22, 22.95, col_w - 0.44, 1.8,
        [
            "3 scénarios : recovery, link flap, multi-fault",
            "3 trafics : none, icmp, iperf_tcp",
            "polling fixe : 2.0 s pour l'expérience principale",
        ],
        font_size=13,
    )

    add_section(slide, x1, 25.7, col_w, 7.65, "DISCUSSION ET LIMITES")
    add_text(
        slide, x1 + 0.22, 26.35, col_w - 0.44, 1.15,
        [
            "Le passage à repeat=10 ne change pas la direction des résultats ;",
            "il renforce leur crédibilité statistique."
        ],
        font_size=13,
    )
    add_bullets(
        slide, x1 + 0.22, 27.8, col_w - 0.44, 2.3,
        [
            "boucle expérimentale complète : données, synthèse et figures",
            "distinction correcte entre panne physique et défaut d'observation",
            "architecture lisible et facilement réplicable pour un TER",
        ],
        font_size=13,
    )
    add_text(slide, x1 + 0.22, 30.45, col_w - 0.44, 0.28, ["Limites :"], font_size=11, color=TEXT_MUTED, bold=True)
    add_text(
        slide, x1 + 0.22, 30.8, col_w - 0.44, 1.65,
        [
            "Mininet + Ryu reste un environnement de simulation.",
            "Les expériences étendues ont moins de répétitions que",
            "la matrice principale et servent surtout à l'analyse de tendance."
        ],
        font_size=11,
        color=TEXT_MUTED,
    )

    # Column 2
    add_section(slide, x2, 6.25, col_w, 6.45, "ARCHITECTURE DU SYSTÈME")
    add_picture(slide, PLOTS / "report" / "architecture.png", x2 + 0.22, 6.88, col_w - 0.44, 3.75)
    add_text(
        slide, x2 + 0.22, 10.85, col_w - 0.44, 0.8,
        [
            "Quatre couches : émulation, observation, noyau jumeau, expérimentation et export.",
            "Ressources observées : switches, liens, hôtes, ports, statistiques et flux."
        ],
        font_size=10,
        color=TEXT_MUTED,
    )

    add_section(slide, x2, 13.05, col_w, 3.95, "MÉCANISME D'ÉTAT ET D'ÉVÉNEMENTS")
    add_text(
        slide, x2 + 0.22, 13.7, col_w - 0.44, 2.4,
        [
            "1. interrogation de l'API Ryu ;",
            "2. construction d'un état structuré ;",
            "3. comparaison avec l'état précédent ;",
            "4. écriture des événements et des exports.",
            "Le jumeau suit donc des transitions, et non une simple photo du réseau."
        ],
        font_size=13,
    )

    add_section(slide, x2, 17.35, col_w, 9.1, "RÉSULTATS PRINCIPAUX : MATRICE REPEAT=10")
    add_picture(slide, PLOTS / "traffic_profile_repeat10_poster.png", x2 + 0.22, 17.98, col_w - 0.44, 3.3)
    add_text(
        slide, x2 + 0.22, 21.45, col_w - 0.44, 0.6,
        ["Tendance la plus stable : le trafic de fond allonge surtout la détection, plus que la reprise."],
        font_size=10,
        color=TEXT_MUTED,
    )
    add_metric(slide, x2 + 0.22, 22.2, 2.9, 1.05, "DETECTION", "0.458 s", "sans trafic")
    add_metric(slide, x2 + 3.42, 22.2, 2.9, 1.05, "DETECTION", "1.144 s", "sous ICMP")
    add_metric(slide, x2 + 6.62, 22.2, 2.9, 1.05, "REPRISE", "≈ 1.87 s", "reste stable", fill=CHIP_ALT)
    add_bullets(
        slide, x2 + 0.22, 23.55, col_w - 0.44, 1.85,
        [
            "`link_flap` et `multi_fault` confirment la capacité à suivre plusieurs transitions",
            "les phases tardives montrent davantage de variance, signe d'un effet de convergence",
        ],
        font_size=12,
    )

    add_section(slide, x2, 26.8, col_w, 3.25, "ENSEIGNEMENTS ROBUSTES")
    add_bullets(
        slide, x2 + 0.22, 27.45, col_w - 0.44, 2.0,
        [
            "la détection panne/reprise est stable dans la matrice principale",
            "l'effet du trafic de fond survit au passage de repeat=3 à repeat=10",
            "l'argument expérimental devient plus convaincant pour un rapport académique",
        ],
        font_size=12.5,
    )

    # Column 3
    add_section(slide, x3, 6.25, col_w, 4.55, "PLAN EXPÉRIMENTAL")
    add_text(
        slide, x3 + 0.22, 6.9, col_w - 0.44, 1.55,
        [
            "Configuration de base : topologie `geant_backbone_8cities_stage3`,",
            "contrôleur `topology_aware_switch`, synchronisation `polling`,",
            "période 2.0 s."
        ],
        font_size=13,
    )
    add_metric(slide, x3 + 0.22, 8.65, 2.85, 0.92, "SCÉNARIOS", "3", "")
    add_metric(slide, x3 + 3.48, 8.65, 2.85, 0.92, "TRAFICS", "3", "")
    add_metric(slide, x3 + 6.74, 8.65, 2.85, 0.92, "RÉPÉTITIONS", "10", "", fill=CHIP_ALT)

    add_section(slide, x3, 11.15, col_w, 6.15, "RÉSULTAT COMPLÉMENTAIRE : POLLING")
    add_picture(slide, PLOTS / "polling_full_matrix_summary.png", x3 + 0.22, 11.8, col_w - 0.44, 2.9)
    add_text(
        slide, x3 + 0.22, 14.88, col_w - 0.44, 0.55,
        ["Conclusion la plus monotone du projet : quand le polling ralentit, la reprise est confirmée plus tard."],
        font_size=10,
        color=TEXT_MUTED,
    )
    add_bullets(
        slide, x3 + 0.22, 15.55, col_w - 0.44, 1.25,
        [
            "1 s  → recovery mean = 0.968 s",
            "2 s  → recovery mean = 1.958 s",
            "4 s  → recovery mean = 3.874 s",
        ],
        font_size=12.5,
    )

    add_section(slide, x3, 17.65, col_w, 4.85, "ROBUSTESSE SÉMANTIQUE ET ÉCHELLE")
    add_bullets(
        slide, x3 + 0.22, 18.3, col_w - 0.44, 1.85,
        [
            "Dégradation d'observation : émission de `SYNC_ERROR`, pas de faux `LINK_DOWN`",
            "Topologie plus complexe : détection 0.628 → 0.725 s",
            "Reprise 1.991 → 2.300 s entre Geant8 et Geant2012_core12",
        ],
        font_size=12,
    )
    add_picture(slide, PLOTS / "p34_full_matrix_summary.png", x3 + 0.22, 20.62, col_w - 0.44, 1.6)

    add_section(slide, x3, 22.85, col_w, 7.2, "CONCLUSION")
    add_text(
        slide, x3 + 0.22, 23.5, col_w - 0.44, 1.45,
        [
            "Le projet valide un PoC de jumeau numérique pour SDN",
            "capable d'observer, structurer, comparer et expliquer",
            "les changements d'état d'un réseau expérimental."
        ],
        font_size=13,
    )
    add_bullets(
        slide, x3 + 0.22, 25.25, col_w - 0.44, 2.65,
        [
            "détection stable des pannes et des reprises",
            "sensibilité nette au trafic de fond",
            "effet monotone du polling sur la reprise",
            "bonne séparation entre faute d'observation et faute physique",
            "base plus solide pour un rapport TER ou une soutenance",
        ],
        font_size=12.5,
    )
    add_text(
        slide, x3 + 0.22, 28.35, col_w - 0.44, 1.1,
        [
            "Prochaine étape raisonnable : étendre les répétitions des expériences complémentaires",
            "et tester des topologies ou contrôleurs plus proches du réel."
        ],
        font_size=11,
        color=TEXT_MUTED,
    )

    footer = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(45.3), Inches(33.11), Inches(1.51))
    footer.fill.solid()
    footer.fill.fore_color.rgb = BLUE_DARK
    footer.line.color.rgb = BLUE_DARK
    add_text(slide, 1.0, 45.62, 16.5, 0.28, ["Source structurante : /docs/report_cn_repeat10.md"], font_size=10, color=RGBColor(215, 228, 241))
    add_text(slide, 1.0, 46.02, 22.0, 0.28, ["Affiche synthétique rédigée en français à partir du rapport de résultats mis à jour."], font_size=10, color=RGBColor(215, 228, 241))

    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_poster()
    print(path)
